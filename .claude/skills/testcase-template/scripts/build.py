#!/usr/bin/env python3
"""
build.py — Sinh file test case Excel từ JSON, giữ nguyên cơ chế của template.

Ranh giới: script CHỈ làm phần cơ khí (clone template, đổ 14 cột đúng vị trí,
chèn section divider, áp style, giữ dropdown, NỚI range công thức Summary, cập
nhật Cover/Summary header, đổ sheet Traceability). Script KHÔNG nghĩ ra test case,
KHÔNG gán priority, KHÔNG dịch — những phần phán đoán đó do agent/Claude làm và
truyền vào qua JSON.

Cách dùng:
    python build.py --input cases.json --template <template.xlsx> --output <out.xlsx>
    python recalc.py <out.xlsx>          # bắt buộc: openpyxl xoá cache công thức

Exit code:
    0 = OK
    1 = validate thất bại (KHÔNG sinh file)
    2 = sinh file được nhưng CÓ VẤN ĐỀ agent phải xử lý tay
        (section vượt sức chứa bảng RESULT BY SECTION, hoặc AC chưa được phủ)

Cấu trúc cases.json: xem reference/cases-json.md (schema đầy đủ + ví dụ chạy được).
Tóm tắt: {"cover": {...}, "acceptance_criteria": [...], "sections": [{"divider", "cases": [...]}]}
"""
import argparse
import json
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

TYPE_ALLOWED = {"UI", "Validation", "Boundary", "Negative", "Functional", "Business rule", "API"}
PRIO_ALLOWED = {"High", "Medium", "Low"}
DATA_START = 6                              # data test case bắt đầu row 6
LAST_COL = 14                               # A..N
SUMMARY_COLS = ["B", "C", "D", "J", "L"]    # các cột Summary tham chiếu qua COUNTIF/COUNTA
ID_RE = re.compile(r"^TC-([A-Z0-9]+)-(\d{3})$")
TRACE_START = 4                             # data sheet Traceability bắt đầu row 4

REQUIRED_COVER = ("module", "version", "source", "create_date")
REQUIRED_EN = ("title_en", "steps_en", "expected_en")
VN_FIELDS = ("title_vn", "steps_vn", "expected_vn")

# ---------------------------------------------------------------- styles
# Style định nghĩa CỨNG ở đây, KHÔNG đọc style của template: nếu kế thừa template
# thì sheet EN mất fill divider (row 6 trống style) còn sheet VN dính fill rác.
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DIVIDER_FILL = PatternFill("solid", fgColor="D9E2F3")
DIVIDER_FONT = Font(name="Arial", size=10, bold=True, color="1F3864")
DATA_FONT = Font(name="Arial", size=10)
NOFILL = PatternFill(fill_type=None)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CENTER_COLS = {1, 3, 4, 10, 11, 12, 13}     # ID, Type, Priority, Result/Bug ID


# ---------------------------------------------------------------- validate
def validate(data):
    """Chặn mọi lỗi có thể sinh ra file 'mở được nhưng sai'. Fail sớm, fail to."""
    errs = []

    cover = data.get("cover") or {}
    for key in REQUIRED_COVER:
        if not str(cover.get(key, "")).strip():
            errs.append(f"cover.{key} bắt buộc (thiếu thì file giữ nguyên metadata mẫu của template)")
    if cover.get("create_date") and not re.match(r"^\d{4}-\d{2}-\d{2}$", str(cover["create_date"])):
        errs.append(f"cover.create_date '{cover['create_date']}' phải dạng YYYY-MM-DD")

    sections = data.get("sections") or []
    if not sections:
        errs.append("Không có section nào trong 'sections'")

    seen_ids = {}
    known_acs = {ac["id"] for ac in data.get("acceptance_criteria", []) if ac.get("id")}

    for si, s in enumerate(sections):
        divider = str(s.get("divider", "")).strip()
        if not divider:
            errs.append(f"sections[{si}]: thiếu 'divider'")
            continue
        m = re.match(r"^([A-Z0-9]+)\.\s+(.+)$", divider)
        if not m:
            errs.append(f"divider '{divider}' phải dạng '<CHỮ CÁI>. <Tên section>', VD 'A. Page & Entry'")
            continue
        letter, sec_name = m.group(1), m.group(2).strip()

        cases = s.get("cases") or []
        if not cases:
            errs.append(f"section '{divider}': không có case nào")

        for c in cases:
            cid = c.get("id", "<thiếu id>")

            im = ID_RE.match(str(cid))
            if not im:
                errs.append(f"{cid}: ID phải dạng TC-<CHỮ CÁI>-<3 số>, VD TC-A-001")
            elif im.group(1) != letter:
                errs.append(f"{cid}: chữ cái ID '{im.group(1)}' không khớp divider '{divider}' "
                            f"(mong đợi TC-{letter}-xxx)")

            if cid in seen_ids:
                errs.append(f"{cid}: TC ID TRÙNG (đã có ở section '{seen_ids[cid]}'). "
                            f"ID trùng làm write_defects.py bỏ sót case fail.")
            else:
                seen_ids[cid] = divider

            if str(c.get("section", "")).strip() != sec_name:
                errs.append(f"{cid}: section '{c.get('section')}' không khớp tên trong divider "
                            f"'{sec_name}' → COUNTIF ở Summary sẽ ra 0")

            if c.get("type") not in TYPE_ALLOWED:
                errs.append(f"{cid}: type '{c.get('type')}' ngoài dropdown {sorted(TYPE_ALLOWED)}")
            if c.get("priority") not in PRIO_ALLOWED:
                errs.append(f"{cid}: priority '{c.get('priority')}' ngoài {sorted(PRIO_ALLOWED)}")

            for f in REQUIRED_EN:
                if not str(c.get(f, "")).strip():
                    errs.append(f"{cid}: thiếu '{f}' (bắt buộc)")
            for f in VN_FIELDS:
                if not str(c.get(f, "")).strip():
                    errs.append(f"{cid}: thiếu '{f}' — sheet VN sẽ có ô trống. "
                                f"Dịch đủ, hoặc chạy lại với --allow-missing-vn.")

            for ac in c.get("acs", []):
                if known_acs and ac not in known_acs:
                    errs.append(f"{cid}: tham chiếu AC '{ac}' không có trong 'acceptance_criteria'")

    return errs


# ---------------------------------------------------------------- ghi data
def clear_data_region(ws, upto=400):
    """Dọn giá trị VÀ style rác từ row 6 xuống, để style của template không rơi
    ngẫu nhiên vào dòng data. Giữ nguyên data validation (dropdown)."""
    end = max(ws.max_row, upto)
    for r in range(DATA_START, end + 1):
        for col in range(1, LAST_COL + 1):
            cell = ws.cell(r, col)
            if isinstance(cell, openpyxl.cell.cell.MergedCell):
                continue
            cell.value = None
            cell.fill = NOFILL
            cell.font = DATA_FONT
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None


def _style_data_row(ws, r):
    for col in range(1, LAST_COL + 1):
        cell = ws.cell(r, col)
        cell.font = DATA_FONT
        cell.border = BORDER
        cell.alignment = CENTER if col in CENTER_COLS else WRAP


def _style_divider_row(ws, r):
    for col in range(1, LAST_COL + 1):
        cell = ws.cell(r, col)
        cell.fill = DIVIDER_FILL
        cell.font = DIVIDER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 22


def write_cases(ws, data, lang, allow_missing_vn=False):
    """Đổ data vào một sheet (EN hoặc VN). Trả về row cuối cùng có data.

    KHÔNG ép row height cho dòng data: để Excel autofit theo wrap_text. Ép cứng
    (VD 28) sẽ cắt cụt Steps / Expected nhiều dòng.
    """
    clear_data_region(ws)
    r = DATA_START
    suffix = "_en" if lang == "en" else "_vn"

    def txt(case, field):
        v = case.get(field + suffix)
        if v in (None, "") and lang == "vn" and allow_missing_vn:
            v = case.get(field + "_en", "")
        return v or ""

    for s in data["sections"]:
        ws.cell(r, 1, s["divider"])
        _style_divider_row(ws, r)
        r += 1
        for c in s["cases"]:
            ws.cell(r, 1, c["id"])                    # A  ID
            ws.cell(r, 2, c["section"])               # B  Section
            ws.cell(r, 3, c["type"])                  # C  Type
            ws.cell(r, 4, c["priority"])              # D  Priority
            ws.cell(r, 5, txt(c, "title"))            # E  Title
            ws.cell(r, 6, txt(c, "precond"))          # F  Precondition
            ws.cell(r, 7, txt(c, "steps"))            # G  Test Steps
            ws.cell(r, 8, txt(c, "data"))             # H  Test Data
            ws.cell(r, 9, txt(c, "expected"))         # I  Expected Result
            ws.cell(r, 10, "Not Run")                 # J  Round 1 Result
            #                                            K  Round 1 Bug ID (trống)
            ws.cell(r, 12, "Not Run")                 # L  Round 2 Result
            #                                            M  Round 2 Bug ID (trống)
            ws.cell(r, 14, txt(c, "note"))            # N  Note / Ref
            _style_data_row(ws, r)
            r += 1
    return r - 1


# ---------------------------------------------------------------- Summary
def find_section_table(ws_summary):
    """Dò bảng RESULT BY SECTION động thay vì hard-code row 21–24.

    Trả về (first_row, capacity). capacity = số dòng liên tiếp ngay dưới header
    có công thức trỏ về sheet 'Test Cases' ở cột B — tức số slot template dựng sẵn.
    """
    header_row = None
    for r in range(1, ws_summary.max_row + 1):
        if str(ws_summary.cell(r, 1).value).strip() == "Section":
            header_row = r
            break
    if header_row is None:
        return None, 0
    first = header_row + 1
    cap = 0
    while True:
        v = ws_summary.cell(first + cap, 2).value
        if isinstance(v, str) and v.startswith("=") and "'Test Cases'" in v:
            cap += 1
        else:
            break
    return first, cap


def update_section_table(ws_summary, data):
    """Ghi tên section thật vào bảng RESULT BY SECTION.

    KHÔNG chèn/xoá dòng: openpyxl insert_rows/delete_rows KHÔNG dịch tham chiếu
    công thức ($A28...) của các bảng bên dưới (BY PRIORITY, LEGEND), nên chèn/xoá
    làm các bảng đó đếm nhầm ô mà file vẫn recalc sạch. Chỉ ghi đè trong slot có
    sẵn; slot thừa thì xoá sạch tên + công thức; thiếu slot thì trả overflow để
    main() thoát với exit code 2.
    """
    first, cap = find_section_table(ws_summary)
    if first is None:
        return {"written": 0, "overflow": [], "capacity": 0,
                "error": "Không tìm thấy bảng RESULT BY SECTION"}

    names = []
    for s in data["sections"]:
        for c in s["cases"]:
            if c["section"] not in names:
                names.append(c["section"])
    overflow = names[cap:] if len(names) > cap else []

    for i in range(cap):
        r = first + i
        if i < len(names):
            ws_summary.cell(r, 1).value = names[i]      # đổi tên; công thức B..J giữ nguyên
        else:
            for col in range(1, 11):                    # slot thừa: xoá tên + công thức
                ws_summary.cell(r, col).value = None
    return {"written": min(len(names), cap), "overflow": overflow, "capacity": cap}


def widen_summary(ws_summary, last_row):
    """Nới mọi range COUNTIF/COUNTA '$X$6:$X$<n>' -> '$X$6:$X$<last_row>'."""
    changed = 0
    for row in ws_summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "'Test Cases'" in cell.value:
                new = cell.value
                for col in SUMMARY_COLS:
                    new = re.sub(rf"\${col}\$6:\${col}\$\d+", f"${col}$6:${col}${last_row}", new)
                if new != cell.value:
                    cell.value = new
                    changed += 1
    return changed


def update_summary_header(ws_summary, cover):
    ws_summary["A2"] = (f"{cover.get('project', 'Fleet Management')}   ·   {cover['module']}"
                        f"   ·   Version {cover['version']}   ·   "
                        f'All figures are live formulas over the "Test Cases" sheet.')


# ---------------------------------------------------------------- headers
def update_en_header(ws_en, cover):
    ws_en["A1"] = f"{cover['module']} — Test Cases"
    ws_en["A2"] = f"Source: {cover['source']}   ·   Version {cover['version']}   ·   {cover['create_date']}"


def update_vn_header(ws_vn, cover):
    ws_vn["A1"] = f"{cover['module']} — Test Cases (Vietnamese reference)"
    ws_vn["A2"] = ('REFERENCE ONLY — mọi số liệu ở sheet Summary được tính từ sheet "Test Cases" (EN).'
                   f"   ·   Source: {cover['source']}   ·   Version {cover['version']}"
                   f"   ·   {cover['create_date']}")


def update_cover(ws_cover, cover):
    ws_cover["A2"] = f"{cover.get('project', 'Fleet Management')} — {cover['module']}"
    ws_cover["B5"] = cover["module"]
    ws_cover["B6"] = cover["version"]
    ws_cover["B7"] = cover["source"]
    ws_cover["E5"] = cover["create_date"]

    # RECORD OF CHANGE: APPEND dòng mới, KHÔNG ghi đè — ghi đè thì làm v2 cho cùng
    # ticket là mất lịch sử thay đổi.
    r = 11
    while True:
        v = ws_cover.cell(r, 1).value
        if v in (None, "") or str(v).startswith("*"):
            break
        r += 1
    ws_cover.cell(r, 1, cover["create_date"])
    ws_cover.cell(r, 2, cover["module"])
    ws_cover.cell(r, 3, cover["version"])
    ws_cover.cell(r, 4, "A" if r == 11 else "M")
    ws_cover.cell(r, 5, cover.get("change_desc", ""))
    ws_cover.cell(r, 6, cover["source"])
    return r


# ---------------------------------------------------------------- traceability
def write_traceability(ws, data):
    """Đổ sheet Traceability: mỗi AC -> danh sách TC phủ nó.

    Đây là lớp bắt sót rẻ nhất của cả harness: AC nào không có TC nào trỏ tới sẽ
    hiện Coverage = MISSING. Trả về danh sách AC chưa được phủ.
    """
    acs = data.get("acceptance_criteria") or []
    for r in range(TRACE_START, max(ws.max_row, TRACE_START) + 1):
        for col in range(1, 6):
            ws.cell(r, col).value = None
    if not acs:
        ws.cell(TRACE_START, 1, "(cases.json không có 'acceptance_criteria' — bỏ qua truy vết)")
        return []

    cover_map = {ac["id"]: [] for ac in acs}
    for s in data["sections"]:
        for c in s["cases"]:
            for ac in c.get("acs", []):
                cover_map.setdefault(ac, []).append(c["id"])

    missing = []
    r = TRACE_START
    for ac in acs:
        tcs = cover_map.get(ac["id"], [])
        if not tcs:
            missing.append(ac["id"])
        ws.cell(r, 1, ac["id"])
        ws.cell(r, 2, ac.get("text", ""))
        ws.cell(r, 3, ", ".join(tcs))
        ws.cell(r, 4, len(tcs))
        ws.cell(r, 5, "MISSING" if not tcs else "Covered")
        for col in range(1, 6):
            cell = ws.cell(r, col)
            cell.font = Font(name="Arial", size=10, bold=not tcs,
                             color="C00000" if not tcs else "000000")
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       horizontal="center" if col in (1, 4, 5) else "left")
        r += 1
    return missing


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--allow-missing-vn", action="store_true",
                    help="Cho phép thiếu bản dịch VN (lấy tạm bản EN). Chỉ dùng khi thật sự cần.")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)

    errs = validate(data)
    if args.allow_missing_vn:
        errs = [e for e in errs if "sheet VN sẽ có ô trống" not in e]
    if errs:
        print("VALIDATION FAILED — không sinh file:", file=sys.stderr)
        for e in errs:
            print("  -", e, file=sys.stderr)
        sys.exit(1)

    cover = data["cover"]
    wb = openpyxl.load_workbook(args.template)

    update_cover(wb["Cover"], cover)
    update_summary_header(wb["Summary"], cover)
    update_en_header(wb["Test Cases"], cover)
    update_vn_header(wb["Test Cases_VN"], cover)

    last_en = write_cases(wb["Test Cases"], data, "en", args.allow_missing_vn)
    last_vn = write_cases(wb["Test Cases_VN"], data, "vn", args.allow_missing_vn)
    last = max(last_en, last_vn)

    sec = update_section_table(wb["Summary"], data)   # TRƯỚC widen_summary
    changed = widen_summary(wb["Summary"], last)
    missing_acs = write_traceability(wb["Traceability"], data)

    wb.save(args.output)

    total_cases = sum(len(s["cases"]) for s in data["sections"])
    out = {
        "output": args.output,
        "total_cases": total_cases,
        "last_data_row": last,
        "sections_written": sec["written"],
        "section_slots_in_template": sec.get("capacity"),
        "summary_formulas_widened": changed,
        "acs_total": len(data.get("acceptance_criteria", [])),
        "acs_missing": missing_acs,
        "next_step": (f"Chạy recalc.py trên {args.output}, rồi kiểm Total == {total_cases} "
                      f"và RESULT BY SECTION cộng đúng"),
    }

    problems = []
    if sec["overflow"]:
        problems.append(f"RESULT BY SECTION chỉ có {sec.get('capacity')} slot nhưng ticket có "
                        f"{sec['written'] + len(sec['overflow'])} section. Section KHÔNG lên bảng: "
                        f"{sec['overflow']}. Summary sẽ không cộng khớp Total — phải thêm dòng tay "
                        f"(và dịch tham chiếu các bảng bên dưới) hoặc gộp section.")
    if missing_acs:
        problems.append(f"{len(missing_acs)} AC chưa có test case nào phủ: {missing_acs}. "
                        f"Xem sheet Traceability, bổ sung case trước khi giao file.")
    if problems:
        out["PROBLEMS"] = problems

    print(json.dumps(out, ensure_ascii=False, indent=2))
    sys.exit(2 if problems else 0)


if __name__ == "__main__":
    main()
